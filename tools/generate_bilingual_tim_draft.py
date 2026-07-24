#!/usr/bin/env python3
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / ".python-packages"))

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "TIM论文/word初稿/A lightweight intensity-assisted framework for semi-solid LiDAR odometry.docx"
OUTPUT = ROOT / "TIM论文/word初稿/A lightweight intensity-assisted framework for semi-solid LiDAR odometry_中英对照.docx"

TRANSLATIONS = [
    "面向半固态激光雷达里程计的轻量级强度辅助框架",
    "摘要：待补充",
    "1. 引言",
    "自动驾驶技术的快速发展推动了三维激光雷达传感器的广泛部署。近年来，高线束半固态激光雷达凭借较低的成本和更高的结构可靠性，逐渐应用于量产车辆。与通常具备 360° 全景视场（Field of View, FoV）和较高测距精度的机械式激光雷达相比，半固态激光雷达受到定向视场（例如水平视场为 120°）和相对较低测量精度的限制。然而，由于此类传感器主要面向车辆前方障碍物检测，其扫描模式十分密集，能够在前向区域提供极高的空间分辨率。",
    "随着半固态激光雷达日益普及，利用此类传感器实现激光雷达里程计与同步定位和建图（Simultaneous Localization and Mapping, SLAM）正受到越来越多关注。然而，将面向机械式激光雷达开发的算法直接应用于半固态激光雷达数据时，往往会出现性能退化。受限视场和更显著的测距噪声使算法难以提取充足且可靠的几何约束，进而不可避免地降低前端里程计的精度与鲁棒性。为了在不过度修改成熟开源框架的前提下实现可靠状态估计，引入一种辅助测量机制具有重要意义。",
    "为应对上述挑战，本文面向高线束半固态激光雷达提出一种轻量级强度辅助插件模块。该方法充分利用此类传感器密集的扫描线，将三维点云投影为高分辨率二维强度图像。通过提取并匹配图像中的光度特征，建立相邻帧之间的稳健数据关联；随后基于匹配特征对应的三维坐标，利用奇异值分解（Singular Value Decomposition, SVD）直接计算帧间变换。在此直接位姿估计基础上，本文进一步将运动补偿嵌入迭代求解回路。通过在迭代过程中执行基于时间戳的运动补偿，算法可同时输出精确的相对位姿和完成去畸变的点云。估计得到的位姿可作为后续点云配准的可靠初始值，从而提升激光雷达里程计的整体鲁棒性与计算效率。本文的主要贡献如下：",
    "本文设计了一条由点云到强度图像的处理流水线，通过空间球面投影，并结合对比度受限自适应直方图均衡化（Contrast-Limited Adaptive Histogram Equalization, CLAHE）和双边滤波，提高信噪比。随后采用两阶段 RANSAC 算法保障特征对应关系的稳健性。",
    "本文提出一种高效的帧间位姿估计方法，将基于时间戳的运动补偿直接嵌入迭代式 SVD 求解过程。该紧耦合策略能够同时实现高精度点云去畸变与相对位姿估计。",
    "本文将所提方法封装为一个极简且高度解耦的插件模块，可无缝集成至现有开源 SLAM 框架。真实场景实验验证了该模块在提升定位精度与鲁棒性方面的有效性。",
    "2. 相关工作",
    "2.1 强度辅助激光雷达里程计",
    "为克服纯几何配准在弱纹理几何结构或非结构化环境中的局限，越来越多的研究开始引入强度测量信息辅助激光雷达里程计。强度信息能够提供独立于空间几何结构的光度观测，为定位提供稳健的视觉线索。借助该信息，相关算法既可以提取具有区分度的视觉特征，也可以构建光度残差以约束位姿估计。Wang 等提出 Intensity-SLAM，通过提取强度特征辅助大规模环境中的定位与建图 [1]。He 等提出 IGICP，融合几何与强度信息以增强点云配准 [2]。类似地，RI-LIO 利用反射率图像提升紧耦合激光雷达惯性里程计性能 [3]；Shan 等则利用成像式激光雷达实现稳健的地点识别 [4]。此外，InTEn-LOAM [5] 及其他强度增强框架 [6] 也表明，引入光度约束能够显著提升数据关联的鲁棒性。然而，这些方法主要针对机械旋转式激光雷达的 360° 均匀扫描模式设计。当其应用于半固态激光雷达时，通常难以充分利用此类传感器前向密集扫描的特性，并可能引入不必要的计算开销。",
    "2.2 面向固态与半固态激光雷达的里程计",
    "随着固态与半固态激光雷达应用日益广泛，面向其特有硬件特性的算法逐渐涌现。由于此类传感器视场受限，在结构化环境中，纯几何配准很容易出现退化。为解决这一问题，Li 等针对固态激光雷达的数据特性，提出了一种面向退化环境的强度增强激光雷达惯性 SLAM 框架 [7]。Pfreundschuh 等提出 COIN-LIO，将光度误差直接纳入迭代配准过程，以补充几何约束 [8]。近期其他工作也探索了面向固态激光雷达的强度增强方法 [9]。尽管这些先进方法取得了较高精度，但通常采用较复杂的紧耦合后端架构，例如迭代卡尔曼滤波器或因子图，以融合高频 IMU 数据。此外，其前端数据关联往往仍依赖计算代价较高的空间迭代最近邻搜索。待补充。",
    "为克服上述方法的局限，本文不再依赖复杂后端和迭代式空间搜索。不同于仅将光度信息作为 ICP 类搜索附加权重的既有方法，本文通过显式强度特征对应关系直接求解帧间变换，从而绕过计算代价较高的迭代搜索过程。此外，本文既不依赖高频 IMU 数据，也不将运动补偿视为孤立的预处理步骤，而是将点级去畸变无缝嵌入迭代式位姿求解。基于这些差异，本文提出一种轻量级、解耦的前端方案，可在避免复杂紧耦合框架高额计算开销的同时，有效缓解半固态激光雷达的固有局限。",
    "3. 方法",
    "本文提出的强度辅助框架被设计为一个轻量级、解耦的插件模块，可无缝集成至现有激光雷达里程计或 SLAM 系统。该模块以连续两帧原始激光雷达点云作为输入，输出两帧之间的六自由度相对位姿以及当前帧去畸变点云。为高效且可靠地实现这一目标，本文方法由三个主要部分构成：（1）点云到强度图像的生成与预处理；（2）两阶段特征提取与匹配；（3）紧耦合点云去畸变与帧间位姿求解。",
    "3.1 强度图像生成与预处理",
    "算法首先将密集三维点云投影为二维强度图像。设一帧原始点云数据为相应符号，其中每个点 pi 均关联一个测得的强度值 Ii。目标是将每个三维点映射至强度图像中的二维像素坐标。",
    "为保证所提框架对不同硬件配置半固态激光雷达的适用性，本文根据传感器规格动态定义生成强度图像的尺寸。设 L 表示激光雷达垂直方向的线束数量。垂直行索引 vi 直接由传感器硬件提供的物理激光通道标签确定，从而生成包含 L 行的图像。对于不提供显式通道标签的传感器，可在柱坐标系中依据俯仰角估计行索引：",
    "公式（1）",
    "其中，θmin 为最小俯仰角，∆θ 为垂直角分辨率。",
    "为确定水平列索引 ui，将 360° 水平方位角均匀划分为 W 个角度区间。为在不同传感器类型下保持一致的特征空间分布纵横比，将水平尺寸 W 定义为垂直线束数量 L 的比例函数，即 W = k×L，其中 k 为可调缩放参数。列索引计算如下：",
    "公式（2）",
    "以 Hesai AT128 为例，该传感器包含 L = 128 条垂直线束。经验设置缩放参数 k = 4，可获得水平分辨率 W = 512，并通过该球面投影模型 [10] 将三维点云有效映射至 128×512 的二维网格。",
    "由于半固态激光雷达在前向区域产生极为密集的扫描，同一个角度区间内可能落入多个点。为解决这种空间混叠问题，本文采用最大强度池化策略：对于每个像素，保留落入对应区间所有点中的最大强度值。该策略能够有效保留交通标志和尖锐结构边缘等高反射率特征，而这些特征具有较强区分度，有利于后续特征提取。",
    "生成初始强度图像后，需要对其进行预处理，以满足特征提取要求。传感器输出的原始强度值被归一化至 [0, 1] 区间。然而，图像对比度通常较低，大部分结构细节被压缩在较窄的低强度区间内，例如 0 至 0.1。为此，本文采用对比度受限自适应直方图均衡化（CLAHE）[11]。与全局直方图均衡化不同，CLAHE 在图像局部区域内执行增强，在提升局部结构细节的同时，避免过度放大均匀暗区中的噪声。",
    "增强后的强度图像仍不可避免地包含高频测量噪声，可能干扰稳定角点特征的提取。为抑制噪声并避免破坏几何边缘，本文在实验中评估了多种常见平滑方法，并发现双边滤波 [12] 的效果最佳。其原因在于双边滤波具备边缘保持特性：不同于对整幅图像进行各向同性模糊的标准高斯平滑，双边滤波同时考虑空间邻近关系与强度差异，可使物体的显著边界保持清晰。",
    "完整的数据处理流程如图 1 所示。图 1(a) 展示半固态激光雷达采集的一帧原始点云。通过球面投影，点云被转换为图 1(b) 所示的原始强度图像。图 1(c) 展示 CLAHE 带来的局部对比度增强效果，图 1(d) 则给出双边滤波后的最终预处理图像，为后续数据关联提供清晰、特征丰富且保持边缘的基础。",
    "待补充",
    "图 1",
    "3.2 光度特征提取与两阶段外点剔除",
    "获得经过预处理且保持边缘的强度图像后，下一步是在连续激光雷达扫描之间建立稳健的数据关联。与传统三维几何特征提取不同，本文直接在二维强度域中提取光度特征。",
    "具体而言，本文从预处理后的强度图像中提取 ORB（Oriented FAST and Rotated BRIEF）特征 [13]。ORB 具有较高计算效率，因此适用于本文场景。光度特征提取结果如图 2(a) 所示，其中上下两个面板分别展示从连续两帧强度图像中提取的 ORB 关键点，并以绿色圆点标记。在分别提取前一帧与当前帧特征后，计算二进制描述子的汉明距离，以获得初始对应关系。随后采用标准最近邻距离比值检验，剔除高度模糊的匹配。",
    "然而，由于激光雷达测量噪声和动态目标的存在，初始匹配集合中不可避免地包含外点。为保证后续位姿估计中数据关联的可靠性，本文提出一种结合二维光度拓扑与三维空间几何的两阶段外点剔除机制。",
    "在第一阶段，本文在二维图像平面内筛选匹配点。将连续强度图像视为连续相机帧，采用随机采样一致性（Random Sample Consensus, RANSAC）算法估计两个视图之间的基础矩阵。该二维 RANSAC 过程施加极线几何约束，从而有效剔除明显错误的光度匹配。经过二维拓扑检验后保留的内点对应关系如图 2(b) 中彩色连线所示。",
    "二维检验能够去除图像级外点，但缺乏空间感知能力，因此仍可能保留三维空间中不合理的对应关系。为此，在第二阶段，根据逆球面投影映射，将剩余二维像素坐标 (u, v) 恢复至其原始三维坐标 (x, y, z)。假设连续扫描的重叠区域主要由静态环境构成，则真实三维对应关系应保持成对空间距离。本文采用三维 RANSAC 算法施加刚体几何一致性约束，并将违反三维刚体变换一致性的匹配视为几何外点予以剔除。如图 2(c) 所示，该步骤进一步净化匹配集合，仅保留物理上有效的几何对应关系。",
    "待补充",
    "图 2",
    "3.3 紧耦合点云去畸变与帧间位姿估计",
    "得到显式对应关系后，可以解析计算由旋转矩阵 R 和平移向量 t 构成的帧间相对位姿。然而，激光雷达以顺序方式采集数据。当传感器持续运动时，同一帧内的点具有不同时间戳，从而在点云 P 中引入显著的运动畸变。对整帧点云统一施加单一刚体变换 (R, t)，将导致次优位姿估计和模糊的建图结果。运动合成 p。",
    "本文不依赖高频 IMU 数据执行孤立的预处理，而是提出一种将点云去畸变直接嵌入迭代式位姿求解回路的紧耦合策略。假设激光雷达在较短的单帧扫描周期内进行匀速运动，则传感器运动可以连续建模。",
]


def get_or_create_style(document, name, base_name, bold=False, color=None, size=10.5):
    styles = document.styles
    if name in styles:
        return styles[name]

    style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = styles[base_name]
    style.font.name = "SimSun"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(size)
    style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor(*color)
    return style


def insert_paragraph_after(paragraph, text, style):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    inserted.style = style
    run = inserted.add_run(text)
    run.font.name = "SimSun"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return inserted


def main():
    document = Document(SOURCE)
    source_paragraphs = list(document.paragraphs)
    if len(source_paragraphs) != len(TRANSLATIONS):
        raise RuntimeError(
            f"Unexpected paragraph count: source={len(source_paragraphs)}, "
            f"translations={len(TRANSLATIONS)}"
        )

    body_style = get_or_create_style(
        document, "Chinese Translation", "Normal", color=(31, 78, 121), size=10.5
    )
    heading_style = get_or_create_style(
        document, "Chinese Translation Heading", "Normal", bold=True, color=(31, 78, 121), size=11
    )
    title_style = get_or_create_style(
        document, "Chinese Translation Title", "Normal", bold=True, color=(31, 78, 121), size=14
    )

    heading_indices = {2, 9, 10, 12, 15, 17, 31, 39}
    for idx, (paragraph, translation) in enumerate(zip(source_paragraphs, TRANSLATIONS)):
        if idx == 0:
            style = title_style
        elif idx in heading_indices:
            style = heading_style
        else:
            style = body_style
            translation = "中文：" + translation
        inserted = insert_paragraph_after(paragraph, translation, style)
        inserted.paragraph_format.space_after = Pt(6)

    document.core_properties.title = "A lightweight intensity-assisted framework for semi-solid LiDAR odometry - bilingual draft"
    document.core_properties.subject = "English-Chinese bilingual draft for review"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
